"""
Adaptador de Excel.

Fornece uma interface para manipulação de arquivos Excel.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from utils.logger import Logger

logger = Logger.get_logger()


class ExcelAdapter:
    """
    Adaptador para manipulação de arquivos Excel.
    
    Attributes
    ----------
    workbook : Workbook
        Instância do workbook do Excel
    """
    
    def __init__(self):
        """Inicializa o adaptador de Excel."""
        self.workbook = Workbook()
        # Remove a planilha padrão criada automaticamente
        default_sheet = self.workbook.active
        self.workbook.remove(default_sheet)
    
    def create_sheet(self, name):
        """
        Cria uma nova planilha no workbook.
        
        Parameters
        ----------
        name : str
            Nome da planilha
            
        Returns
        -------
        worksheet
            A planilha criada
        """
        logger.debug(f"Criando planilha: {name}")
        return self.workbook.create_sheet(name)
    
    def save(self, filename):
        """
        Salva o workbook em um arquivo.
        
        Parameters
        ----------
        filename : str
            Nome do arquivo
            
        Returns
        -------
        bool
            True se o salvamento foi bem-sucedido, False caso contrário
        """
        try:
            # Certifica-se de que o diretório existe
            os.makedirs(os.path.dirname(os.path.abspath(filename)), exist_ok=True)
            logger.info(f"Salvando arquivo Excel: {filename}")
            self.workbook.save(filename)
            return True
        except Exception as e:
            logger.error(f"Erro ao salvar arquivo Excel: {str(e)}", exc_info=True)
            Logger.capture_error(e)
            return False
    
    @staticmethod
    def create_style(bold=False, size=11, name="Calibri", color="000000", 
                     bg_color=None, horizontal="general", vertical="bottom", 
                     border=None, wrap_text=False):
        """
        Cria um estilo para células do Excel.
        
        Parameters
        ----------
        bold : bool, optional
            Se a fonte deve ser negrito
        size : int, optional
            Tamanho da fonte
        name : str, optional
            Nome da fonte
        color : str, optional
            Cor da fonte em formato hexadecimal
        bg_color : str, optional
            Cor de fundo em formato hexadecimal
        horizontal : str, optional
            Alinhamento horizontal (center, left, right, general)
        vertical : str, optional
            Alinhamento vertical (top, center, bottom)
        border : dict, optional
            Configuração de bordas
        wrap_text : bool, optional
            Se o texto deve ser quebrado automaticamente
            
        Returns
        -------
        dict
            Dicionário com as configurações de estilo
        """
        style = {
            'font': Font(bold=bold, size=size, name=name, color=color),
            'alignment': Alignment(
                horizontal=horizontal, 
                vertical=vertical,
                wrap_text=wrap_text
            )
        }
        
        if bg_color:
            style['fill'] = PatternFill(fill_type="solid", fgColor=bg_color)
        
        if border:
            borders = {}
            for side, border_style in border.items():
                if border_style:
                    borders[side] = Side(border_style=border_style)
            
            style['border'] = Border(
                left=borders.get('left', Side(border_style=None)),
                right=borders.get('right', Side(border_style=None)),
                top=borders.get('top', Side(border_style=None)),
                bottom=borders.get('bottom', Side(border_style=None))
            )
        
        return style
    
    @staticmethod
    def apply_style(cell, style):
        """
        Aplica um estilo a uma célula.
        
        Parameters
        ----------
        cell : Cell
            Célula a receber o estilo
        style : dict
            Dicionário com as configurações de estilo
        """
        if 'font' in style:
            cell.font = style['font']
        
        if 'alignment' in style:
            cell.alignment = style['alignment']
        
        if 'fill' in style:
            cell.fill = style['fill']
        
        if 'border' in style:
            cell.border = style['border']
    
    @staticmethod
    def get_column_letter(col_idx):
        """
        Converte um índice de coluna em letra.
        
        Parameters
        ----------
        col_idx : int
            Índice da coluna (1-based)
            
        Returns
        -------
        str
            Letra da coluna (A, B, C, ..., Z, AA, AB, ...)
        """
        from openpyxl.utils import get_column_letter
        return get_column_letter(col_idx) 